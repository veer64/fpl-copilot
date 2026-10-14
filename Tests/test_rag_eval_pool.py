"""RAG eval part 1 (2026-10-14): the question set's mix and the pure helpers of eval/rag_eval/build_pool.py --
de-duplication across retrieval variants, the seeded shuffle, the as_of check, and the production-shaped ranking
(recency, per-article cap, player gate). No database, no network."""
import json
import sys
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[1]
for p in (str(REPO), str(REPO / "squad"), str(REPO / "eval")):
    if p not in sys.path:
        sys.path.insert(0, p)

import news_search as ns                                   # noqa: E402
from rag_eval import build_pool as bp                      # noqa: E402

UTC = timezone.utc
QUESTIONS = REPO / "eval" / "rag_eval" / "questions_v1.json"
REQUIRED_MIX = {"player_current_news": 9, "club_news": 4, "name_variant": 3, "conflict": 3, "stale": 3,
                "no_news": 4, "ambiguous": 2, "vague": 2}


def _cand(cid, item, score, fetched, published=None, source="club", players=None, player_ids=None, element_id=None,
          date_source="meta"):
    return {"id": cid, "news_item_id": item, "score": score, "r_dense": None, "r_kw": None, "source": source, "club": None,
            "url": None, "published_at": published, "fetched_at": fetched, "date_source": date_source, "headline": f"h{cid}",
            "element_id": element_id, "chunk_text": f"text {cid}", "players": players or [], "player_ids": player_ids or []}


# ---- the question set -------------------------------------------------------------------------------------------

def test_questions_v1_has_exactly_the_required_mix():
    doc = json.loads(QUESTIONS.read_text(encoding="utf-8"))
    qs = doc["questions"]
    assert len(qs) == 30
    assert Counter(q["situation"] for q in qs) == REQUIRED_MIX
    assert doc["situations"] == REQUIRED_MIX
    assert len({q["qid"] for q in qs}) == 30
    for q in qs:
        for k in ("qid", "situation", "question", "players", "clubs", "as_of", "expected_outcome", "right_behaviour"):
            assert q.get(k) is not None, (q["qid"], k)
        as_of = bp.parse_as_of(q["as_of"])
        assert as_of.tzinfo is not None
        assert datetime(2026, 10, 3, tzinfo=UTC) <= as_of <= datetime(2026, 10, 13, 23, 59, tzinfo=UTC)
    # mostly 2026-10-03 .. 2026-10-10 09:00Z, a few on 2026-10-13
    late = [q for q in qs if bp.parse_as_of(q["as_of"]) > datetime(2026, 10, 10, 9, 0, tzinfo=UTC)]
    assert all(bp.parse_as_of(q["as_of"]).date() == datetime(2026, 10, 13).date() for q in late)
    assert len(late) <= 6


def test_questions_v1_at_most_one_question_per_player_outside_conflict_and_stale():
    qs = json.loads(QUESTIONS.read_text(encoding="utf-8"))["questions"]
    seen = Counter()
    for q in qs:
        if q["situation"] in ("conflict", "stale"):
            continue
        for p in q["players"]:
            seen[p.lower()] += 1
    assert all(n == 1 for n in seen.values()), seen


def test_questions_without_players_or_clubs_are_the_vague_ones_and_clubs_only_the_club_ones():
    qs = json.loads(QUESTIONS.read_text(encoding="utf-8"))["questions"]
    for q in qs:
        if q["situation"] == "vague":
            assert q["players"] == [] and q["clubs"] == []
        elif q["situation"] == "club_news":
            assert q["players"] == [] and len(q["clubs"]) == 1
        else:
            assert len(q["players"]) == 1 and q["clubs"] == []


# ---- de-duplication across variants --------------------------------------------------------------------------------

def test_merge_variants_dedups_by_chunk_and_keeps_every_rank():
    t = datetime(2026, 10, 8, tzinfo=UTC)
    a, b, c = _cand(1, 10, 0.9, t), _cand(2, 11, 0.8, t), _cand(3, 12, 0.7, t)
    merged = bp.merge_variants({"dense_only": [a, b], "keyword_only": [b, c], "hybrid_no_recency": [c, a, b],
                                "hybrid_production": [a]})
    assert sorted(m["chunk_id"] for m in merged) == [1, 2, 3]
    by = {m["chunk_id"]: m["found_by"] for m in merged}
    assert by[1] == {"dense_only": 1, "hybrid_no_recency": 2, "hybrid_production": 1}
    assert by[2] == {"dense_only": 2, "keyword_only": 1, "hybrid_no_recency": 3}
    assert by[3] == {"keyword_only": 2, "hybrid_no_recency": 1}
    assert all("score" not in m and "boosted" not in m for m in merged)       # no score reaches the labeller


def test_merge_variants_empty():
    assert bp.merge_variants({v: [] for v in bp.VARIANTS}) == []


# ---- the seeded shuffle ----------------------------------------------------------------------------------------------

def test_shuffle_is_a_seeded_permutation():
    rows = [{"chunk_id": i} for i in range(12)]
    s1 = bp.shuffle_candidates(rows, bp.SHUFFLE_SEED, "q07")
    s2 = bp.shuffle_candidates(rows, bp.SHUFFLE_SEED, "q07")
    assert s1 == s2
    assert sorted(r["chunk_id"] for r in s1) == list(range(12))
    assert s1 is not rows and [r["chunk_id"] for r in rows] == list(range(12))       # the input is untouched
    other = bp.shuffle_candidates(rows, bp.SHUFFLE_SEED, "q08")
    assert [r["chunk_id"] for r in other] != [r["chunk_id"] for r in s1]
    assert isinstance(bp.SHUFFLE_SEED, int)


# ---- as_of ------------------------------------------------------------------------------------------------------------

def test_check_as_of_reports_only_chunks_fetched_after_as_of():
    as_of = datetime(2026, 10, 9, 9, 0, tzinfo=UTC)
    ok = _cand(1, 10, 0.9, as_of - timedelta(hours=1))
    edge = _cand(2, 11, 0.9, as_of)
    late = _cand(3, 12, 0.9, as_of + timedelta(seconds=1))
    assert bp.check_as_of([ok, edge], as_of) == []
    assert [v["chunk_id"] for v in bp.check_as_of([ok, late, edge], as_of)] == [3]


def test_production_sql_filters_live_items_by_as_of():
    sql = ns._sql(True, 2)
    assert "n.fetched_at <= %(as_of)s" in sql


def test_parse_as_of_reads_utc_iso():
    assert bp.parse_as_of("2026-10-09T09:00:00Z") == datetime(2026, 10, 9, 9, 0, tzinfo=UTC)
    with pytest.raises(ValueError):
        bp.parse_as_of("2026-10-09 09:00")


# ---- the production-shaped ranking -------------------------------------------------------------------------------------

def test_rank_candidates_gate_keeps_only_about_rows_when_players_given():
    as_of = datetime(2026, 10, 9, 9, 0, tzinfo=UTC)
    t = as_of - timedelta(days=1)
    about = _cand(1, 10, 0.5, t, published=t, player_ids=[379])
    filler = _cand(2, 11, 0.9, t, published=t, player_ids=[999])
    fpl_about = _cand(3, 12, 0.4, t, published=t, source="fpl", element_id=379, date_source="source_field")
    out = bp.rank_candidates([about, filler, fpl_about], as_of, requested_ids=[379], players_given=True, recency=False, top=5)
    assert [c["id"] for c in out] == [1, 3]
    assert all(c["about"] for c in out)
    out2 = bp.rank_candidates([about, filler, fpl_about], as_of, requested_ids=[], players_given=False, recency=False, top=5)
    assert [c["id"] for c in out2] == [2, 1, 3] and not any(c["about"] for c in out2)


def test_rank_candidates_caps_two_chunks_per_article_and_tops_n():
    as_of = datetime(2026, 10, 9, 9, 0, tzinfo=UTC)
    t = as_of - timedelta(days=1)
    cands = [_cand(i, 10, 1.0 - i / 100, t, published=t) for i in range(1, 5)] + [_cand(9, 11, 0.5, t, published=t)]
    out = bp.rank_candidates(cands, as_of, requested_ids=[], players_given=False, recency=False, top=5)
    assert [c["id"] for c in out] == [1, 2, 9]
    assert [c["id"] for c in bp.rank_candidates(cands, as_of, [], False, recency=False, top=2)] == [1, 2]


def test_rank_candidates_recency_reorders_like_production():
    as_of = datetime(2026, 10, 9, 9, 0, tzinfo=UTC)
    old = _cand(1, 10, 0.60, as_of - timedelta(days=9), published=as_of - timedelta(days=9))
    new = _cand(2, 11, 0.40, as_of - timedelta(hours=2), published=as_of - timedelta(hours=2))
    no_recency = bp.rank_candidates([old, new], as_of, [], False, recency=False, top=5)
    with_recency = bp.rank_candidates([old, new], as_of, [], False, recency=True, top=5)
    assert [c["id"] for c in no_recency] == [1, 2]
    assert [c["id"] for c in with_recency] == [2, 1]
    # the boosted score is production's: score * 0.5 ** (age_days / half life)
    boosted = {c["id"]: c["boosted"] for c in with_recency}
    assert boosted[1] == pytest.approx(0.60 * ns.recency_multiplier(9.0))


# ---- pool v2 (2026-10-14): the ungated dense variant and the reuse of stored query embeddings -------------------------

def test_variants_keep_the_four_production_ones_first_and_add_dense_ungated_last():
    assert bp.VARIANTS[:4] == ("dense_only", "keyword_only", "hybrid_no_recency", "hybrid_production")
    assert bp.VARIANTS[-1] == "dense_ungated"


def test_gate_active_only_when_a_named_player_resolved():
    assert bp.gate_active(["Isak"], [379]) is True
    assert bp.gate_active(["Palmer"], []) is False            # named but ambiguous / unmapped: nothing gated
    assert bp.gate_active([], []) is False                    # clubs or a vague question: no gate
    assert bp.gate_active([], [379]) is False                 # ids without names cannot happen in production


def test_dense_ungated_ignores_the_gate_and_the_recency_multiplier():
    as_of = datetime(2026, 10, 9, 9, 0, tzinfo=UTC)
    t_old, t_new = as_of - timedelta(days=9), as_of - timedelta(hours=2)
    about = _cand(1, 10, 0.5, t_new, published=t_new, player_ids=[379])
    filler_old = _cand(2, 11, 0.9, t_old, published=t_old, player_ids=[999])
    out = bp.dense_ungated([about, filler_old], as_of, requested_ids=[379])
    assert [c["id"] for c in out] == [2, 1]                   # the filler stays, raw vector order, the old one not demoted
    assert out[0]["about"] is False and out[1]["about"] is True
    assert all(c["boosted"] == pytest.approx(c["score"]) for c in out)


def test_dense_ungated_keeps_the_two_chunks_per_article_cap_and_the_top():
    as_of = datetime(2026, 10, 9, 9, 0, tzinfo=UTC)
    t = as_of - timedelta(days=1)
    cands = [_cand(i, 10, 1.0 - i / 100, t, published=t) for i in range(1, 5)] + [_cand(9, 11, 0.5, t, published=t)]
    assert [c["id"] for c in bp.dense_ungated(cands, as_of, requested_ids=[379])] == [1, 2, 9]
    many = [_cand(i, 100 + i, 1.0 - i / 100, t, published=t) for i in range(1, 9)]
    assert len(bp.dense_ungated(many, as_of, requested_ids=[379])) == bp.TOP_PER_VARIANT


def test_dense_ungated_refuses_a_chunk_fetched_after_as_of():
    as_of = datetime(2026, 10, 9, 9, 0, tzinfo=UTC)
    ok = _cand(1, 10, 0.9, as_of - timedelta(hours=1))
    late = _cand(3, 12, 0.9, as_of + timedelta(seconds=1))
    with pytest.raises(ValueError):
        bp.dense_ungated([ok, late], as_of, requested_ids=[379])


def test_dense_ungated_is_skipped_when_no_player_is_gated():
    as_of = datetime(2026, 10, 9, 9, 0, tzinfo=UTC)
    c = _cand(1, 10, 0.9, as_of - timedelta(hours=1))
    assert bp.dense_ungated([c], as_of, requested_ids=[]) is None


def test_only_by_names_the_rows_a_single_variant_found():
    t = datetime(2026, 10, 8, tzinfo=UTC)
    a, b, c = _cand(1, 10, 0.9, t), _cand(2, 11, 0.8, t), _cand(3, 12, 0.7, t)
    merged = bp.merge_variants({"dense_only": [a], "hybrid_production": [a, c], "dense_ungated": [a, b, c]})
    assert bp.only_by(merged, "dense_ungated") == [2]
    assert bp.only_by(merged, "dense_only") == []


def test_stored_query_embeddings_are_reused_without_an_embedding_call(monkeypatch):
    import embeddings as em
    pool = {"questions": [{"qid": "q01", "query_embedding": [0.1, 0.2]}, {"qid": "q02", "query_embedding": [0.3, 0.4]}]}
    assert bp.embeddings_from_pool(pool) == {"q01": [0.1, 0.2], "q02": [0.3, 0.4]}
    assert bp.embeddings_from_pool({"questions": [{"qid": "q01"}]}) == {}

    def boom(*a, **k):
        raise AssertionError("embeddings.embed must not be called when every question has a stored vector")
    monkeypatch.setattr(em, "embed", boom)
    questions = [{"qid": "q01", "question": "a"}, {"qid": "q02", "question": "b"}]
    stats = {}
    out = bp.embed_questions(questions, conn=None, stats=stats, reuse=bp.embeddings_from_pool(pool))
    assert out == {"q01": [0.1, 0.2], "q02": [0.3, 0.4]} and stats.get("reused") == 2
    with pytest.raises(AssertionError):                       # one vector missing -> the real embed is needed
        bp.embed_questions(questions + [{"qid": "q03", "question": "c"}], conn=None, stats={}, reuse=bp.embeddings_from_pool(pool))


def test_label_row_count_rule():
    """one outcome row for an ambiguous / no-news question (or any question with no candidates), plus every candidate"""
    assert bp.label_row_count("ambiguous", 0) == 1
    assert bp.label_row_count("no_news", 0) == 1
    assert bp.label_row_count("no_news", 5) == 6                # the tool's verdict row, then the ungated candidates
    assert bp.label_row_count("player_current_news", 0) == 1
    assert bp.label_row_count("player_current_news", 4) == 4


def test_workbook_shows_the_outcome_row_then_the_candidates_for_a_no_news_question(tmp_path):
    import openpyxl
    t = "2026-10-08T12:00:00Z"
    doc = {"questions": [
        {"qid": "q24", "situation": "no_news", "question": "any news on marmoush?", "players": ["Marmoush"], "clubs": [],
         "as_of": t, "expected_outcome": "no_news_for", "right_behaviour": "say none"},
        {"qid": "q27", "situation": "ambiguous", "question": "is palmer fit?", "players": ["Palmer"], "clubs": [],
         "as_of": t, "expected_outcome": "ambiguous", "right_behaviour": "ask"}]}

    def cand(cid):
        return {"chunk_id": cid, "item_id": 100 + cid, "source": "fpl", "club": None, "headline": f"h{cid}",
                "published_at": "2026-10-07T10:00:00+00:00", "fetched_at": "2026-10-07T10:30:00+00:00",
                "date_source": "source_field", "chunk_text": f"text {cid}", "found_by": {"dense_ungated": cid}, "about": False}
    pool = {"questions": [
        {"qid": "q24", "situation": "no_news", "n_candidates": 2, "candidates": [cand(1), cand(2)],
         "tool_outcome": {"search_mode": "hybrid", "n_results": 0, "no_news_for": ["Marmoush"], "ambiguous_players": {}}},
        {"qid": "q27", "situation": "ambiguous", "n_candidates": 0, "candidates": [],
         "tool_outcome": {"search_mode": "dense_only", "n_results": 0, "no_news_for": [], "ambiguous_players": {"Palmer": ["a", "b"]}}}]}
    sha = bp.write_workbook(doc, pool, tmp_path / "w.xlsx")
    assert len(sha) == 12
    ws = openpyxl.load_workbook(tmp_path / "w.xlsx")["label"]
    assert ws["A1"].value == bp.LABEL_RULE
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    q24 = [r for r in rows if r[0] == "q24"]
    q27 = [r for r in rows if r[0] == "q27"]
    assert [r[3] for r in q24] == ["tool", 1, 2] and q24[0][7] == bp.OUTCOME_QUESTION
    assert sorted(r[8] for r in q24[1:]) == ["text 1", "text 2"]
    assert [r[3] for r in q27] == ["tool"]
    header = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    assert not any(k in str(h).lower() for h in header for k in ("variant", "rank", "score"))


def test_rank_candidates_date_basis_follows_production():
    """an item whose date_source is not a publisher claim ages from fetched_at (news_search.date_basis)"""
    as_of = datetime(2026, 10, 9, 9, 0, tzinfo=UTC)
    c = _cand(1, 10, 0.5, as_of - timedelta(days=2), published=as_of - timedelta(days=30), date_source="first_seen")
    out = bp.rank_candidates([c], as_of, [], False, recency=True, top=5)
    assert out[0]["boosted"] == pytest.approx(0.5 * ns.recency_multiplier(2.0))
