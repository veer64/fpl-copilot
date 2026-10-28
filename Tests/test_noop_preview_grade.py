"""The no-op preview definition (Prompt E2, 2026-10-28; eval/reports/confirm_lock_v2.md Part 0).

For a request whose requested setting is ALREADY TRUE on the active squad, the preview step of the
must-call chain is satisfied by EITHER a set_my_squad(confirm=false) preview OR a reply that says the
setting is already in place, with NO write in that turn. Grammar: the pseudo-tool `already(<word>)` as a
slot alternative, used per row (v4_10: get_my_xi:6 only). Still FAIL: any write; ignoring the request;
a claim that something changed. The validation set below is the task's, written before the grader change.
No network, no database: every test builds its own calls and replies (the workbook test reads the two
committed spec files)."""
import inspect
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
for p in (str(REPO), str(REPO / "eval"), str(REPO / "squad")):
    if p not in sys.path:
        sys.path.insert(0, p)

import run_agent_evals as rae                             # noqa: E402

OLD = "get_my_xi > set_my_squad(confirm=false)"
NEW = "get_my_xi > set_my_squad(confirm=false)|already(captain)"
ALREADY = "Just to be clear -- Haaland is already your recorded captain (set as of GW6), so that part doesn't need changing."


def calls(*spec):
    out = []
    for s in spec:
        if isinstance(s, tuple):
            out.append({"tool": s[0], "args": s[1], "error": None})
        else:
            out.append({"tool": s, "args": {}, "error": None})
    return out


def grade(chain, cs, reply="", wrote=False):
    q = {"must_call": chain, "must_not": "set_my_squad(confirm=true)"}
    return rae.grade(q, cs, setup_calls=[], reply=reply, wrote=wrote)


# ---------------------------------------------------------------- the validation set
def test_preview_call_passes():
    assert grade(NEW, calls("get_my_xi", "set_my_squad"), reply="Here's the preview. Shall I confirm?") == ("PASS", "")
    assert grade(NEW, calls("get_my_xi", ("set_my_squad", {"confirm": False}))) == ("PASS", "")


def test_already_in_place_with_no_write_passes():
    assert grade(NEW, calls("get_my_xi"), reply=ALREADY, wrote=False) == ("PASS", "")


def test_write_without_preview_fails():
    res, reason = grade(NEW, calls("get_my_xi", ("set_my_squad", {"confirm": True})), reply="Done, Haaland is your captain.")
    assert res == "FAIL" and reason


def test_already_then_a_write_fails():
    # the reply says "already" but the turn wrote: (a) fails, whichever way the write shows
    res, _ = grade(NEW, calls("get_my_xi", ("set_my_squad", {"confirm": True})), reply=ALREADY)
    assert res == "FAIL"
    res, reason = grade(NEW, calls("get_my_xi"), reply=ALREADY, wrote=True)
    assert res == "FAIL" and "missing" in reason


def test_an_answer_about_something_else_fails():
    res, reason = grade(NEW, calls("get_my_xi"), reply="Arsenal's next three fixtures are Liverpool, Hull City and Newcastle.")
    assert res == "FAIL" and "missing" in reason and "already(captain)" in reason


def test_a_claim_that_something_changed_fails():
    for reply in ("I've set Haaland as your captain.", "Haaland is now your captain for GW9.",
                  "Your captain has been updated to Haaland.", "Haaland is already your captain. I've applied the captain change."):
        res, _ = grade(NEW, calls("get_my_xi"), reply=reply)
        assert res == "FAIL", reply
    # a negated or no-change sentence is not a claim
    ok = "Haaland is already your captain -- no change needed, I have not changed anything."
    assert grade(NEW, calls("get_my_xi"), reply=ok) == ("PASS", "")


# ---------------------------------------------------------------- the grammar
def test_the_v4_8_chain_is_unchanged_by_the_definition():
    res, reason = grade(OLD, calls("get_my_xi"), reply=ALREADY)
    assert res == "FAIL" and "missing set_my_squad(confirm=false) after step 1" == reason


def test_already_needs_the_word_and_already_in_one_sentence():
    assert grade(NEW, calls("get_my_xi"), reply="ALREADY your captain: Haaland.") == ("PASS", "")
    res, _ = grade(NEW, calls("get_my_xi"), reply="He is your captain. I already told you the XI.")
    assert res == "FAIL"
    res, _ = grade(NEW, calls("get_my_xi"), reply="Haaland is already your vice.")
    assert res == "FAIL"


def test_already_requires_the_earlier_steps():
    res, reason = grade(NEW, calls(), reply=ALREADY)        # get_my_xi never called
    assert res == "FAIL" and reason.startswith("missing get_my_xi")


def test_already_is_never_a_tool_call():
    assert rae.parse_tool_spec("already(captain)") == ("already", ("word", "captain"))
    assert rae.call_matches({"tool": "already", "args": {"word": "captain"}}, "already(captain)") is False
    assert rae.grade_must_not("already(captain)", calls("get_my_xi")) == (True, "")


def test_grade_defaults_keep_the_old_signature():
    q = {"must_call": OLD, "must_not": "set_my_squad(confirm=true)"}
    assert rae.grade(q, calls("get_my_xi", "set_my_squad")) == ("PASS", "")
    assert rae.grade(q, calls("get_my_xi"))[0] == "FAIL"


def test_run_one_hands_the_reply_and_the_write_flag_to_grade():
    src = inspect.getsource(rae.run_one)
    assert 'reply=qturn["reply"]' in src and 'wrote=rec.get("wrote_squad_version"' in src


# ---------------------------------------------------------------- the spec version
def test_v4_10_is_v4_8_plus_the_one_cell():
    import openpyxl
    d = REPO / "eval" / "agent_evals"
    a = openpyxl.load_workbook(d / "tool_questions_v4_8.xlsx", read_only=True)
    b = openpyxl.load_workbook(d / "tool_questions_v4_10.xlsx", read_only=True)
    assert a.sheetnames == b.sheetnames
    diffs = []
    for name in a.sheetnames:
        ra = list(a[name].iter_rows(values_only=True))
        rb = list(b[name].iter_rows(values_only=True))
        if name == "changes":
            assert rb[:len(ra)] == ra and len(rb) == len(ra) + 1
            assert rb[-1][0] == "get_my_xi#6" and rb[-1][1] == "must call"
            continue
        assert len(ra) == len(rb), name
        for i, (x, y) in enumerate(zip(ra, rb)):
            if x != y:
                diffs.append((name, i, x, y))
    assert len(diffs) == 1
    name, i, x, y = diffs[0]
    assert name == "get_my_xi" and x[0] == 6 and y[0] == 6
    assert x[3] == OLD and y[3] == NEW
    assert [v for j, v in enumerate(x) if j != 3] == [v for j, v in enumerate(y) if j != 3]
